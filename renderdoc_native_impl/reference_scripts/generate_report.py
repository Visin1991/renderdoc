"""Generate a complete analysis report from RenderDoc capture data.

This script combines:
  1. Excel report generation (from capture_analysis.json)
  2. Cluster lighting visualization (heatmap, per-pixel lights, comparison)
  3. Embedding the comparison image into the Excel workbook

By default, generates Excel report only. The comparison image is generated
temporarily, embedded into the Excel workbook, and then cleaned up.
No files are left in the visualization/ directory.

Use --debug to generate and keep all visualization images:
  - capture_analysis_perpixel_lights.png           (per-pixel light count heatmap)
  - capture_analysis_backbuffer.png                (backbuffer export)
  - capture_analysis_comparison.png                (side-by-side comparison)
  - capture_analysis_cluster_lighting.png          (3D grid collapsed heatmap)
  - capture_analysis_cluster_lighting_2d_tiles.png (2D tile heatmap)
  - capture_analysis_cluster_lighting_z_slices.png (per-Z-slice breakdown)
  - capture_analysis_perpixel_lights_zslice.png    (Z-slice assignment per pixel)

Usage:
    python generate_report.py [capture_dir] [--debug] [--output name.xlsx]

Examples:
    python generate_report.py                          # Use ./capture/ folder
    python generate_report.py capture --debug          # With debug images
    python generate_report.py 楼船4月1 --output report.xlsx

Dependencies:
    pip install openpyxl matplotlib numpy pillow
"""
'''
# Default: generate Excel report only (comparison image embedded, no files kept)
python generate_report.py

# Debug mode: generate Excel + keep all visualization images
python generate_report.py --debug

# Specify capture directory and output filename
python generate_report.py 楼船4月1 --output report.xlsx

# Only generate visualizations (debug mode), skip Excel
python generate_report.py --no-excel --debug

# Only generate Excel, skip visualization entirely (no comparison image embedded)
python generate_report.py --no-vis
'''

import json
import sys
import os
import argparse
import importlib

# ── Ensure script directory is on path for sibling imports ──
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)


def find_capture_dir(user_input=None):
    """Resolve the capture directory from user input or default locations."""
    if user_input:
        candidate = os.path.abspath(user_input)
        if os.path.isdir(candidate):
            return candidate
        # Maybe user passed a JSON file path
        if os.path.isfile(candidate) and candidate.endswith(".json"):
            return os.path.dirname(candidate)
        print(f"Error: '{user_input}' is not a valid directory or JSON file.")
        sys.exit(1)

    # Default: look for ./capture/ subfolder
    default = os.path.join(SCRIPT_DIR, "capture")
    if os.path.isdir(default):
        return default

    # Fallback: script directory itself
    if os.path.exists(os.path.join(SCRIPT_DIR, "capture_analysis.json")):
        return SCRIPT_DIR

    print("Error: No capture directory found.")
    print("  Expected: ./capture/capture_analysis.json")
    print(f"  Or specify a directory: python {os.path.basename(__file__)} <capture_dir>")
    sys.exit(1)


def load_json(capture_dir):
    """Load capture_analysis.json from the capture directory."""
    json_path = os.path.join(capture_dir, "capture_analysis.json")
    if not os.path.exists(json_path):
        print(f"Error: {json_path} not found.")
        sys.exit(1)
    print(f"Reading: {json_path}")
    with open(json_path, "r", encoding="utf-8") as f:
        return json.load(f), json_path


# ═══════════════════════════════════════════════════════════════════════════════
# Part 1: Excel Report Generation
# ═══════════════════════════════════════════════════════════════════════════════

def generate_excel(data, json_path, output_xlsx, comparison_img_path=None):
    """Generate Excel workbook with all analysis sheets + optional embedded image."""
    # Import json_to_excel module
    import json_to_excel as j2e

    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XlImage
        from openpyxl.styles import Font
    except ImportError:
        print("Error: openpyxl is required. Install it with: pip install openpyxl")
        sys.exit(1)

    wb = Workbook()
    wb.remove(wb.active)

    # Write all standard sheets
    sheets = [
        ("概览",           j2e.write_overview_sheet),
        ("材质统计",       j2e.write_material_stats_sheet),
        ("模型统计",       j2e.write_mesh_stats_sheet),
        ("材质x模型组合",  j2e.write_material_mesh_combo_sheet),
        ("子模块统计",     j2e.write_submodule_stats_sheet),
        ("原始数据",       j2e.write_raw_drawcalls_sheet),
        ("Dispatches",     j2e.write_dispatches_sheet),
        ("Textures",       j2e.write_textures_sheet),
        ("Buffers",        j2e.write_buffers_sheet),
        ("纹理格式分布",   j2e.write_texture_format_sheet),
        ("纹理分辨率分布", j2e.write_texture_resolution_sheet),
    ]

    for name, func in sheets:
        print(f"  Writing [{name}] ...")
        func(wb, data)

    # ── Embed comparison image into a new sheet ──
    if comparison_img_path and os.path.exists(comparison_img_path):
        print(f"  Writing [可视化对比] ...")
        ws_img = wb.create_sheet("可视化对比")
        ws_img.sheet_properties.tabColor = "FF0066"

        ws_img.cell(row=1, column=1, value="Cluster Lighting vs Final Rendering").font = Font(
            bold=True, size=14, color="4472C4"
        )
        ws_img.cell(row=2, column=1, value="Left: Per-pixel light count heatmap  |  Right: Final rendered frame").font = Font(
            size=10, color="666666"
        )

        try:
            img = XlImage(comparison_img_path)
            # Scale image to fit nicely in the sheet
            # openpyxl uses EMU internally; we set width/height in pixels
            max_width = 1200  # pixels
            aspect = img.height / img.width if img.width > 0 else 1
            if img.width > max_width:
                img.width = max_width
                img.height = int(max_width * aspect)
            ws_img.add_image(img, "A4")
            print(f"    Embedded: {comparison_img_path}")
        except Exception as e:
            print(f"    Warning: Failed to embed image: {e}")
            ws_img.cell(row=4, column=1, value=f"[Image embed failed: {e}]")
    else:
        if comparison_img_path:
            print(f"  Skipping [可视化对比] - image not found: {comparison_img_path}")

    # Save workbook
    print(f"\nSaving Excel: {output_xlsx}")
    wb.save(output_xlsx)
    print(f"  {len(wb.sheetnames)} sheets generated:")
    for i, name in enumerate(wb.sheetnames, 1):
        print(f"    {i:2d}. {name}")


# ═══════════════════════════════════════════════════════════════════════════════
# Part 2: Visualization Generation
# ═══════════════════════════════════════════════════════════════════════════════

def generate_visualizations(data, capture_dir):
    """Generate all visualization images (only called in --debug mode).

    Returns the path to the comparison image (or None if not generated).

    Generated images:
      - perpixel_lights.png           (per-pixel light count heatmap)
      - backbuffer.png                (backbuffer export)
      - comparison.png                (side-by-side comparison)
      - cluster_lighting.png          (3D grid collapsed heatmap)
      - cluster_lighting_2d_tiles.png (2D tile sub-figure)
      - cluster_lighting_z_slices.png (per-Z-slice breakdown)
      - perpixel_lights_zslice.png    (Z-slice assignment per pixel)
    """
    # Import visualization module
    import visualize_cluster_lighting as vcl

    cluster_data = data.get("clusterLighting")
    if cluster_data is None:
        print("\n  ClusterLighting data not found in JSON, skipping visualizations.")
        return None

    base_name = "capture_analysis"
    vis_dir = os.path.join(capture_dir, "visualization")
    os.makedirs(vis_dir, exist_ok=True)

    title_prefix = os.path.basename(capture_dir) if capture_dir else ""

    # Sanitize data
    cluster_data = vcl._sanitize_cluster_data(cluster_data)
    cell_counts = cluster_data.get("cellLightCounts", [])

    comparison_path = None

    # ── Full cluster lighting heatmap (3D grid collapsed) ──
    if cell_counts:
        print("\n  Generating cluster lighting heatmap...")
        cl_output = os.path.join(vis_dir, f"{base_name}_cluster_lighting.png")
        vcl.create_cluster_lighting_heatmap(cluster_data, cl_output, title_prefix)

    # ── Per-pixel light count ──
    scene_depth_meta = data.get("sceneDepth")
    light_grid_params = data.get("lightGridParams")
    depth_bin_path = os.path.join(capture_dir, "scene_depth.bin")
    depth_output = os.path.join(vis_dir, f"{base_name}_perpixel_lights.png")

    perpixel_ok = False
    if cell_counts and scene_depth_meta and os.path.exists(depth_bin_path):
        print("\n  Generating per-pixel light count heatmap...")
        perpixel_ok = vcl.create_perpixel_light_heatmap(
            cluster_data, scene_depth_meta, light_grid_params,
            depth_bin_path, depth_output, title_prefix,
        )
    else:
        missing = []
        if not cell_counts:
            missing.append("cellLightCounts")
        if not scene_depth_meta:
            missing.append("sceneDepth")
        if not os.path.exists(depth_bin_path):
            missing.append("scene_depth.bin")
        print(f"\n  Skipping per-pixel visualization (missing: {', '.join(missing)}).")

    # ── BackBuffer export + Comparison ──
    backbuffer_meta = data.get("backbuffer")
    bb_bin_path = os.path.join(capture_dir, "backbuffer.bin")
    bb_direct_png = os.path.join(capture_dir, "backbuffer_direct.png")
    bb_png_path = os.path.join(vis_dir, f"{base_name}_backbuffer.png")

    bb_ready = False
    if backbuffer_meta:
        import shutil
        if backbuffer_meta.get("savedViaSaveTexture") and os.path.exists(bb_direct_png):
            print("\n  Using SaveTexture-generated BackBuffer PNG...")
            try:
                shutil.copy2(bb_direct_png, bb_png_path)
                print(f"    Copied -> {bb_png_path}")
                bb_ready = True
            except Exception as e:
                print(f"    Failed to copy: {e}")

        if not bb_ready and os.path.exists(bb_bin_path):
            print("\n  Exporting BackBuffer from raw binary...")
            bb_ready = vcl.export_backbuffer_png(backbuffer_meta, bb_bin_path, bb_png_path)

        if bb_ready and perpixel_ok and os.path.exists(depth_output):
            comparison_output = os.path.join(vis_dir, f"{base_name}_comparison.png")
            print("\n  Generating comparison image (Heatmap vs BackBuffer)...")
            vcl.create_comparison_image(depth_output, bb_png_path, comparison_output, title_prefix)
            comparison_path = comparison_output
        elif not bb_ready:
            print("\n  BackBuffer export failed, skipping comparison image.")
        elif not perpixel_ok:
            print("\n  Per-pixel heatmap not available, skipping comparison image.")
    else:
        print("\n  BackBuffer metadata not found, skipping comparison image.")

    return comparison_path


def generate_comparison_only(data, capture_dir):
    """Generate only the comparison image (temporarily) for Excel embedding.

    This is used in non-debug mode: we create the comparison image,
    embed it into Excel, then the caller cleans up the visualization/ dir.

    Returns the path to the comparison image (or None if not generated).
    """
    import visualize_cluster_lighting as vcl

    cluster_data = data.get("clusterLighting")
    if cluster_data is None:
        print("\n  ClusterLighting data not found in JSON, skipping comparison.")
        return None

    base_name = "capture_analysis"
    vis_dir = os.path.join(capture_dir, "visualization")
    os.makedirs(vis_dir, exist_ok=True)

    title_prefix = os.path.basename(capture_dir) if capture_dir else ""

    # Sanitize data
    cluster_data = vcl._sanitize_cluster_data(cluster_data)
    cell_counts = cluster_data.get("cellLightCounts", [])

    # ── Per-pixel light count heatmap (needed for comparison) ──
    scene_depth_meta = data.get("sceneDepth")
    light_grid_params = data.get("lightGridParams")
    depth_bin_path = os.path.join(capture_dir, "scene_depth.bin")
    depth_output = os.path.join(vis_dir, f"{base_name}_perpixel_lights.png")

    perpixel_ok = False
    if cell_counts and scene_depth_meta and os.path.exists(depth_bin_path):
        print("\n  Generating per-pixel light count heatmap (temporary)...")
        perpixel_ok = vcl.create_perpixel_light_heatmap(
            cluster_data, scene_depth_meta, light_grid_params,
            depth_bin_path, depth_output, title_prefix,
        )
    else:
        missing = []
        if not cell_counts:
            missing.append("cellLightCounts")
        if not scene_depth_meta:
            missing.append("sceneDepth")
        if not os.path.exists(depth_bin_path):
            missing.append("scene_depth.bin")
        print(f"\n  Skipping comparison (missing: {', '.join(missing)}).")
        return None

    if not perpixel_ok:
        print("\n  Per-pixel heatmap generation failed, skipping comparison.")
        return None

    # ── BackBuffer ──
    backbuffer_meta = data.get("backbuffer")
    bb_bin_path = os.path.join(capture_dir, "backbuffer.bin")
    bb_direct_png = os.path.join(capture_dir, "backbuffer_direct.png")
    bb_png_path = os.path.join(vis_dir, f"{base_name}_backbuffer.png")

    bb_ready = False
    if backbuffer_meta:
        import shutil
        if backbuffer_meta.get("savedViaSaveTexture") and os.path.exists(bb_direct_png):
            print("\n  Using SaveTexture-generated BackBuffer PNG...")
            try:
                shutil.copy2(bb_direct_png, bb_png_path)
                bb_ready = True
            except Exception as e:
                print(f"    Failed to copy: {e}")

        if not bb_ready and os.path.exists(bb_bin_path):
            print("\n  Exporting BackBuffer from raw binary...")
            bb_ready = vcl.export_backbuffer_png(backbuffer_meta, bb_bin_path, bb_png_path)
    else:
        print("\n  BackBuffer metadata not found, skipping comparison.")
        return None

    if not bb_ready:
        print("\n  BackBuffer export failed, skipping comparison.")
        return None

    # ── Comparison image ──
    comparison_output = os.path.join(vis_dir, f"{base_name}_comparison.png")
    print("\n  Generating comparison image (temporary for Excel embedding)...")
    vcl.create_comparison_image(depth_output, bb_png_path, comparison_output, title_prefix)

    if os.path.exists(comparison_output):
        return comparison_output
    return None


def cleanup_visualization_dir(capture_dir):
    """Remove all files in the visualization/ directory and the directory itself.

    Called after embedding images into Excel in non-debug mode so that
    no visualization files are left on disk.
    """
    vis_dir = os.path.join(capture_dir, "visualization")
    if not os.path.isdir(vis_dir):
        return
    removed = 0
    for fname in os.listdir(vis_dir):
        fpath = os.path.join(vis_dir, fname)
        if os.path.isfile(fpath):
            os.remove(fpath)
            removed += 1
    # Remove the directory itself if empty
    try:
        os.rmdir(vis_dir)
    except OSError:
        pass  # directory not empty (unexpected files), leave it
    if removed:
        print(f"  Cleaned up {removed} temporary visualization file(s).")


# ═══════════════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="Generate analysis report (Excel + Visualizations) from RenderDoc capture data.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python generate_report.py                          # Use ./capture/ folder
  python generate_report.py capture --debug          # With debug visualizations
  python generate_report.py 楼船4月1 --output report.xlsx
        """,
    )
    parser.add_argument(
        "capture_dir", nargs="?", default=None,
        help="Path to the capture directory (default: ./capture/)",
    )
    parser.add_argument(
        "--debug", action="store_true",
        help="Generate and keep all visualization images in visualization/ folder",
    )
    parser.add_argument(
        "--output", "-o", default=None,
        help="Output Excel filename (default: capture_analysis.xlsx in capture dir)",
    )
    parser.add_argument(
        "--no-excel", action="store_true",
        help="Skip Excel generation, only generate visualizations",
    )
    parser.add_argument(
        "--no-vis", action="store_true",
        help="Skip visualization generation, only generate Excel",
    )

    args = parser.parse_args()

    # Resolve capture directory
    capture_dir = find_capture_dir(args.capture_dir)
    print(f"Capture directory: {capture_dir}")
    print(f"Debug mode: {'ON' if args.debug else 'OFF'}")
    print("=" * 60)

    # Load JSON data
    data, json_path = load_json(capture_dir)

    comparison_path = None

    # ── Step 1: Generate visualizations ──
    if args.no_vis:
        print("\n[Step 1] Skipping visualizations (--no-vis)")
    elif args.debug:
        # Debug mode: generate ALL images and keep them
        print("\n[Step 1] Generating all visualizations (debug mode)...")
        comparison_path = generate_visualizations(data, capture_dir)
        if comparison_path:
            print(f"\n  Comparison image: {comparison_path}")
    else:
        # Default mode: generate comparison image temporarily for Excel embedding
        print("\n[Step 1] Generating comparison image for Excel embedding...")
        comparison_path = generate_comparison_only(data, capture_dir)
        if comparison_path:
            print(f"\n  Comparison image (temporary): {comparison_path}")

    # ── Step 2: Generate Excel report ──
    output_xlsx = None
    if not args.no_excel:
        print("\n[Step 2] Generating Excel report...")
        if args.output:
            output_xlsx = args.output
            if not os.path.isabs(output_xlsx):
                output_xlsx = os.path.join(capture_dir, output_xlsx)
        else:
            output_xlsx = os.path.join(capture_dir, "capture_analysis.xlsx")

        generate_excel(data, json_path, output_xlsx, comparison_path)
    else:
        print("\n[Step 2] Skipping Excel generation (--no-excel)")

    # ── Step 3: Cleanup temporary files in non-debug mode ──
    if not args.debug and not args.no_vis:
        print("\n[Step 3] Cleaning up temporary visualization files...")
        cleanup_visualization_dir(capture_dir)

    # ── Summary ──
    print("\n" + "=" * 60)
    print("Report generation complete!")
    if args.debug:
        vis_dir = os.path.join(capture_dir, "visualization")
        if os.path.isdir(vis_dir):
            vis_files = [f for f in os.listdir(vis_dir) if f.endswith(".png")]
            if vis_files:
                print(f"\nVisualization images ({len(vis_files)}):")
                for f in sorted(vis_files):
                    print(f"  - visualization/{f}")
    if output_xlsx:
        print(f"\nExcel report: {output_xlsx}")
    print("=" * 60)


if __name__ == "__main__":
    main()
