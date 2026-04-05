"""Generate a complete analysis report from RenderDoc capture data.

Enhanced version that includes Overdraw statistics in the Excel report.

This script combines:
  1. Excel report generation (from capture_analysis.json) — with Overdraw data
  2. Cluster lighting visualization (heatmap, per-pixel lights, comparison)
  3. Embedding the comparison image into the Excel workbook

Usage:
    python generate_report.py [capture_dir] [--debug] [--output name.xlsx]

Examples:
    python generate_report.py                          # Use ./capture/ folder
    python generate_report.py capture --debug          # With debug images
    python generate_report.py 楼船4月1 --output report.xlsx

Dependencies:
    pip install openpyxl matplotlib numpy pillow
"""

import json
import sys
import os
import argparse
import importlib

# ── Ensure script directories are on path ──
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
REFERENCE_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "renderdoc_native_impl", "reference_scripts")

for d in [SCRIPT_DIR, REFERENCE_DIR]:
    if d not in sys.path:
        sys.path.insert(0, d)


def find_capture_dir(user_input=None):
    """Resolve the capture directory from user input or default locations."""
    if user_input:
        candidate = os.path.abspath(user_input)
        if os.path.isdir(candidate):
            return candidate
        if os.path.isfile(candidate) and candidate.endswith(".json"):
            return os.path.dirname(candidate)
        print(f"Error: '{user_input}' is not a valid directory or JSON file.")
        sys.exit(1)

    default = os.path.join(SCRIPT_DIR, "capture")
    if os.path.isdir(default):
        return default

    if os.path.exists(os.path.join(SCRIPT_DIR, "capture_analysis.json")):
        return SCRIPT_DIR

    print("Error: No capture directory found.")
    print("  Expected: ./capture/capture_analysis.json")
    sys.exit(1)


def load_json(capture_dir):
    """Load capture_analysis.json from the capture directory."""
    json_path = os.path.join(capture_dir, "capture_analysis.json")
    if not os.path.exists(json_path):
        print(f"Error: {json_path} not found.")
        sys.exit(1)
    print(f"Reading: {json_path}")
    with open(json_path, "r", encoding="utf-8") as f:
        return json.load(f), json_path


# ═══════════════════════════════════════════════════════════════════════════════
# Part 1: Excel Report Generation (Enhanced with Overdraw)
# ═══════════════════════════════════════════════════════════════════════════════

def generate_excel(data, json_path, output_xlsx, comparison_img_path=None):
    """Generate Excel workbook with all analysis sheets + Overdraw + optional image."""
    # Import the enhanced json_to_excel from scripts/
    import json_to_excel as j2e

    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XlImage
        from openpyxl.styles import Font
    except ImportError:
        print("Error: openpyxl is required. Install it with: pip install openpyxl")
        sys.exit(1)

    has_od = j2e._has_overdraw_data(data)
    print(f"  Overdraw data available: {has_od}")

    wb = Workbook()
    wb.remove(wb.active)

    # Write all standard sheets (enhanced with Overdraw)
    sheets = [
        ("概览",           j2e.write_overview_sheet),
        ("材质统计",       j2e.write_material_stats_sheet),
        ("模型统计",       j2e.write_mesh_stats_sheet),
        ("材质x模型组合",  j2e.write_material_mesh_combo_sheet),
        ("子模块统计",     j2e.write_submodule_stats_sheet),
        ("原始数据",       j2e.write_raw_drawcalls_sheet),
        ("Dispatches",     j2e.write_dispatches_sheet),
        ("Textures",       j2e.write_textures_sheet),
        ("Buffers",        j2e.write_buffers_sheet),
        ("纹理格式分布",   j2e.write_texture_format_sheet),
        ("纹理分辨率分布", j2e.write_texture_resolution_sheet),
    ]

    # Add Overdraw sheet if data is available
    if has_od:
        sheets.append(("Overdraw分析", j2e.write_overdraw_analysis_sheet))

    for name, func in sheets:
        print(f"  Writing [{name}] ...")
        func(wb, data)

    # ── Embed comparison image ──
    if comparison_img_path and os.path.exists(comparison_img_path):
        print(f"  Writing [可视化对比] ...")
        ws_img = wb.create_sheet("可视化对比")
        ws_img.sheet_properties.tabColor = "FF0066"

        ws_img.cell(row=1, column=1, value="Cluster Lighting vs Final Rendering").font = Font(
            bold=True, size=14, color="4472C4"
        )
        ws_img.cell(row=2, column=1, value="Left: Per-pixel light count heatmap  |  Right: Final rendered frame").font = Font(
            size=10, color="666666"
        )

        try:
            img = XlImage(comparison_img_path)
            max_width = 1200
            aspect = img.height / img.width if img.width > 0 else 1
            if img.width > max_width:
                img.width = max_width
                img.height = int(max_width * aspect)
            ws_img.add_image(img, "A4")
        except Exception as e:
            ws_img.cell(row=4, column=1, value=f"[Image embed failed: {e}]")

    # Save workbook
    print(f"\nSaving Excel: {output_xlsx}")
    wb.save(output_xlsx)
    print(f"  {len(wb.sheetnames)} sheets generated:")
    for i, name in enumerate(wb.sheetnames, 1):
        print(f"    {i:2d}. {name}")


# ═══════════════════════════════════════════════════════════════════════════════
# Part 2: Visualization Generation (delegates to reference_scripts)
# ═══════════════════════════════════════════════════════════════════════════════

def generate_visualizations(data, capture_dir):
    """Generate all visualization images (only called in --debug mode)."""
    try:
        import visualize_cluster_lighting as vcl
    except ImportError:
        print("  [WARN] visualize_cluster_lighting not available, skipping visualizations.")
        return None

    cluster_data = data.get("clusterLighting")
    if cluster_data is None:
        print("\n  ClusterLighting data not found in JSON, skipping visualizations.")
        return None

    base_name = "capture_analysis"
    vis_dir = os.path.join(capture_dir, "visualization")
    os.makedirs(vis_dir, exist_ok=True)

    title_prefix = os.path.basename(capture_dir) if capture_dir else ""
    cluster_data = vcl._sanitize_cluster_data(cluster_data)
    cell_counts = cluster_data.get("cellLightCounts", [])
    comparison_path = None

    if cell_counts:
        print("\n  Generating cluster lighting heatmap...")
        cl_output = os.path.join(vis_dir, f"{base_name}_cluster_lighting.png")
        vcl.create_cluster_lighting_heatmap(cluster_data, cl_output, title_prefix)

    scene_depth_meta = data.get("sceneDepth")
    light_grid_params = data.get("lightGridParams")
    depth_bin_path = os.path.join(capture_dir, "scene_depth.bin")
    depth_output = os.path.join(vis_dir, f"{base_name}_perpixel_lights.png")

    perpixel_ok = False
    if cell_counts and scene_depth_meta and os.path.exists(depth_bin_path):
        print("\n  Generating per-pixel light count heatmap...")
        perpixel_ok = vcl.create_perpixel_light_heatmap(
            cluster_data, scene_depth_meta, light_grid_params,
            depth_bin_path, depth_output, title_prefix,
        )

    backbuffer_meta = data.get("backbuffer")
    bb_bin_path = os.path.join(capture_dir, "backbuffer.bin")
    bb_direct_png = os.path.join(capture_dir, "backbuffer_direct.png")
    bb_png_path = os.path.join(vis_dir, f"{base_name}_backbuffer.png")

    bb_ready = False
    if backbuffer_meta:
        import shutil
        if backbuffer_meta.get("savedViaSaveTexture") and os.path.exists(bb_direct_png):
            try:
                shutil.copy2(bb_direct_png, bb_png_path)
                bb_ready = True
            except:
                pass
        if not bb_ready and os.path.exists(bb_bin_path):
            bb_ready = vcl.export_backbuffer_png(backbuffer_meta, bb_bin_path, bb_png_path)

        if bb_ready and perpixel_ok and os.path.exists(depth_output):
            comparison_output = os.path.join(vis_dir, f"{base_name}_comparison.png")
            vcl.create_comparison_image(depth_output, bb_png_path, comparison_output, title_prefix)
            comparison_path = comparison_output

    return comparison_path


def generate_comparison_only(data, capture_dir):
    """Generate only the comparison image (temporarily) for Excel embedding."""
    try:
        import visualize_cluster_lighting as vcl
    except ImportError:
        return None

    cluster_data = data.get("clusterLighting")
    if cluster_data is None:
        return None

    base_name = "capture_analysis"
    vis_dir = os.path.join(capture_dir, "visualization")
    os.makedirs(vis_dir, exist_ok=True)

    title_prefix = os.path.basename(capture_dir) if capture_dir else ""
    cluster_data = vcl._sanitize_cluster_data(cluster_data)
    cell_counts = cluster_data.get("cellLightCounts", [])

    scene_depth_meta = data.get("sceneDepth")
    light_grid_params = data.get("lightGridParams")
    depth_bin_path = os.path.join(capture_dir, "scene_depth.bin")
    depth_output = os.path.join(vis_dir, f"{base_name}_perpixel_lights.png")

    perpixel_ok = False
    if cell_counts and scene_depth_meta and os.path.exists(depth_bin_path):
        perpixel_ok = vcl.create_perpixel_light_heatmap(
            cluster_data, scene_depth_meta, light_grid_params,
            depth_bin_path, depth_output, title_prefix,
        )

    if not perpixel_ok:
        return None

    backbuffer_meta = data.get("backbuffer")
    bb_bin_path = os.path.join(capture_dir, "backbuffer.bin")
    bb_direct_png = os.path.join(capture_dir, "backbuffer_direct.png")
    bb_png_path = os.path.join(vis_dir, f"{base_name}_backbuffer.png")

    bb_ready = False
    if backbuffer_meta:
        import shutil
        if backbuffer_meta.get("savedViaSaveTexture") and os.path.exists(bb_direct_png):
            try:
                shutil.copy2(bb_direct_png, bb_png_path)
                bb_ready = True
            except:
                pass
        if not bb_ready and os.path.exists(bb_bin_path):
            bb_ready = vcl.export_backbuffer_png(backbuffer_meta, bb_bin_path, bb_png_path)

    if not bb_ready:
        return None

    comparison_output = os.path.join(vis_dir, f"{base_name}_comparison.png")
    vcl.create_comparison_image(depth_output, bb_png_path, comparison_output, title_prefix)

    return comparison_output if os.path.exists(comparison_output) else None


def cleanup_visualization_dir(capture_dir):
    """Remove temporary visualization files."""
    vis_dir = os.path.join(capture_dir, "visualization")
    if not os.path.isdir(vis_dir):
        return
    removed = 0
    for fname in os.listdir(vis_dir):
        fpath = os.path.join(vis_dir, fname)
        if os.path.isfile(fpath):
            os.remove(fpath)
            removed += 1
    try:
        os.rmdir(vis_dir)
    except OSError:
        pass
    if removed:
        print(f"  Cleaned up {removed} temporary visualization file(s).")


# ═══════════════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="Generate analysis report (Excel + Visualizations + Overdraw) from RenderDoc capture data.",
    )
    parser.add_argument("capture_dir", nargs="?", default=None,
                        help="Path to the capture directory (default: ./capture/)")
    parser.add_argument("--debug", action="store_true",
                        help="Generate and keep all visualization images")
    parser.add_argument("--output", "-o", default=None,
                        help="Output Excel filename")
    parser.add_argument("--no-excel", action="store_true",
                        help="Skip Excel generation")
    parser.add_argument("--no-vis", action="store_true",
                        help="Skip visualization generation")

    args = parser.parse_args()

    capture_dir = find_capture_dir(args.capture_dir)
    print(f"Capture directory: {capture_dir}")
    print(f"Debug mode: {'ON' if args.debug else 'OFF'}")
    print("=" * 60)

    data, json_path = load_json(capture_dir)

    # Check for overdraw data
    od_stats = data.get("overdrawStats")
    if od_stats and od_stats.get("available"):
        print(f"Overdraw data: AVAILABLE (Overall: {data.get('overview', {}).get('overallOverdraw', 0):.4f}x)")
    else:
        print("Overdraw data: NOT AVAILABLE")

    comparison_path = None

    # Step 1: Visualizations
    if args.no_vis:
        print("\n[Step 1] Skipping visualizations (--no-vis)")
    elif args.debug:
        print("\n[Step 1] Generating all visualizations (debug mode)...")
        comparison_path = generate_visualizations(data, capture_dir)
    else:
        print("\n[Step 1] Generating comparison image for Excel embedding...")
        comparison_path = generate_comparison_only(data, capture_dir)

    # Step 2: Excel report
    output_xlsx = None
    if not args.no_excel:
        print("\n[Step 2] Generating Excel report (with Overdraw)...")
        if args.output:
            output_xlsx = args.output
            if not os.path.isabs(output_xlsx):
                output_xlsx = os.path.join(capture_dir, output_xlsx)
        else:
            output_xlsx = os.path.join(capture_dir, "capture_analysis.xlsx")

        generate_excel(data, json_path, output_xlsx, comparison_path)

    # Step 3: Cleanup
    if not args.debug and not args.no_vis:
        print("\n[Step 3] Cleaning up temporary files...")
        cleanup_visualization_dir(capture_dir)

    # Summary
    print("\n" + "=" * 60)
    print("Report generation complete!")
    if output_xlsx:
        print(f"\nExcel report: {output_xlsx}")
    print("=" * 60)


if __name__ == "__main__":
    main()
